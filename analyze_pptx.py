import sys
import io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

from pptx import Presentation
from pptx.util import Inches, Pt, Emu

prs = Presentation('Sentinel_AI_Final_Presentation.pptx')
print(f"Total Slides: {len(prs.slides)}")
print(f"Slide Width: {prs.slide_width} EMU = {prs.slide_width / 914400:.2f} inches")
print(f"Slide Height: {prs.slide_height} EMU = {prs.slide_height / 914400:.2f} inches")

for i, slide in enumerate(prs.slides):
    print(f"\n{'='*80}")
    print(f"SLIDE {i+1}")
    print(f"{'='*80}")
    layout = slide.slide_layout
    print(f"Layout: {layout.name}")
    
    for s in slide.shapes:
        print(f"\n  Shape: {s.name}")
        print(f"  Type: {s.shape_type}")
        print(f"  Position: left={s.left}, top={s.top}, width={s.width}, height={s.height}")
        
        if s.has_text_frame:
            for p_idx, para in enumerate(s.text_frame.paragraphs):
                text = para.text
                if text.strip():
                    align = para.alignment
                    font_info = []
                    for run in para.runs:
                        try:
                            color_val = run.font.color.rgb if run.font.color and run.font.color.rgb else 'None'
                        except:
                            color_val = 'theme/inherited'
                        fi = f"font={run.font.name}, size={run.font.size}, bold={run.font.bold}, italic={run.font.italic}, color={color_val}"
                        font_info.append(fi)
                    # Replace emojis for safe printing
                    safe_text = text.encode('ascii', 'replace').decode('ascii')[:120]
                    print(f"  Para {p_idx}: '{safe_text}' | align={align} | {'; '.join(font_info)}")
        
        if hasattr(s, 'image'):
            try:
                print(f"  IMAGE: {s.image.content_type}, size={len(s.image.blob)} bytes")
            except:
                pass

# Also extract slide backgrounds
print("\n\n" + "="*80)
print("SLIDE BACKGROUNDS")
print("="*80)
for i, slide in enumerate(prs.slides):
    bg = slide.background
    fill = bg.fill
    print(f"\nSlide {i+1}: fill type = {fill.type}")
