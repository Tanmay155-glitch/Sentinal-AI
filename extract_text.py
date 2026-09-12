import sys
import io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

from pptx import Presentation
from pptx.util import Inches, Pt, Emu

prs = Presentation('Sentinel_AI_Final_Presentation.pptx')

# Get background colors
for i, slide in enumerate(prs.slides):
    bg = slide.background
    fill = bg.fill
    try:
        if fill.type is not None:
            fc = fill.fore_color
            print(f"Slide {i+1} bg color: {fc.rgb}")
    except:
        print(f"Slide {i+1} bg: theme/inherited")

# Get full text content per slide
print("\n\n=== FULL TEXT CONTENT ===")
for i, slide in enumerate(prs.slides):
    print(f"\n--- SLIDE {i+1} ---")
    for s in slide.shapes:
        if s.has_text_frame:
            full_text = s.text_frame.text.strip()
            if full_text:
                print(full_text)
                print("---")
